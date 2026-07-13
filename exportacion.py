from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _ajustar_hoja(ws) -> None:
    verde = "167D73"
    claro = "DCEFED"
    borde = Side(style="thin", color="C8D4D3")

    for celda in ws[1]:
        celda.fill = PatternFill("solid", fgColor=verde)
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = Border(bottom=borde)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for columna in ws.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(len(str(c.value or "")) for c in columna) + 2
        ws.column_dimensions[letra].width = min(max(ancho, 10), 34)
        for celda in columna:
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    for fila in range(2, ws.max_row + 1):
        if fila % 2 == 0:
            for celda in ws[fila]:
                celda.fill = PatternFill("solid", fgColor=claro)


def exportar_excel(resumen: pd.DataFrame, detalle: pd.DataFrame, fichadas: pd.DataFrame) -> bytes:
    detalle_export = detalle.drop(columns=[c for c in detalle.columns if c.endswith(" min")], errors="ignore")
    observaciones = detalle_export[detalle_export["Observaciones"].astype(str).str.strip() != ""]

    salida = BytesIO()
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen RRHH", index=False)
        detalle_export.to_excel(writer, sheet_name="Detalle Jornadas", index=False)
        observaciones.to_excel(writer, sheet_name="Observaciones", index=False)
        fichadas.to_excel(writer, sheet_name="Fichadas originales", index=False)

    salida.seek(0)
    wb = load_workbook(salida)
    for ws in wb.worksheets:
        _ajustar_hoja(ws)
    wb["Resumen RRHH"].sheet_view.showGridLines = False
    wb["Resumen RRHH"].row_dimensions[1].height = 26

    final = BytesIO()
    wb.save(final)
    return final.getvalue()
